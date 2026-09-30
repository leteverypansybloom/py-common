"""Unit tests for SharePointSource (Microsoft Graph) adapter.

Every Graph call goes to an in-memory fake (FakeGraph) that returns
canned responses shaped like real Microsoft Graph responses, so no
test contacts SharePoint or needs credentials. Workbooks are
synthetic, built with openpyxl at test time.

Authentication is not tested here beyond "the adapter asks the
injected token provider and sends a bearer token": which credential
flow supplies the token is an IT decision still to be made.

RAP Principles:
- Reproducible: Fixed fake responses; no network, clock or sleep
- Auditable: Each test name states the behaviour it locks in
- Transparent: No real tenants, sites or secrets anywhere
"""

from __future__ import annotations

import logging
from collections.abc import Mapping
from io import BytesIO
from typing import Any
from unittest.mock import MagicMock, Mock

import pytest
import requests
from openpyxl import Workbook

from py_common.adapters.sharepoint import (
    GRAPH_URL,
    SharePointConfig,
    SharePointSource,
)
from py_common.config import ConfigError
from py_common.contract import Column, Contract, Worksheet
from py_common.errors import NotFound, ServiceError, VersionMismatch
from py_common.model import Outcome, SourceItem, digest
from py_common.pipeline import Pipeline
from py_common.ports import Source

TOKEN = "fake-token-not-a-secret"  # pragma: allowlist secret
SITE_ID = "contoso.sharepoint.com,site-guid,web-guid"
DRIVE_ID = "b!drive-id"
SITE_URL = f"{GRAPH_URL}/sites/contoso.sharepoint.com:/sites/DataSite"
DRIVES_URL = f"{GRAPH_URL}/sites/{SITE_ID}/drives"
CHILDREN_URL = (
    f"{GRAPH_URL}/drives/{DRIVE_ID}/root:/Incoming%20Files/2026:/children"
)


# --------------------------------------------------------------------
# Fake Microsoft Graph
# --------------------------------------------------------------------


class FakeResponse:
    """Minimal stand-in for requests.Response."""

    def __init__(
        self,
        status_code: int = 200,
        body: Any = None,
        content: bytes = b"",
        headers: Mapping[str, str] | None = None,
    ) -> None:
        """Build a canned response.

        Args:
            status_code: HTTP status code.
            body: Parsed JSON body; None means "not JSON".
            content: Raw bytes (for file downloads).
            headers: Response headers.
        """
        self.status_code = status_code
        self._body = body
        self.content = content
        self.headers = dict(headers or {})

    def json(self) -> Any:
        """Return the JSON body, or raise like requests does.

        Returns:
            Parsed body.

        Raises:
            ValueError: Body is not JSON.
        """
        if self._body is None:
            raise ValueError("not JSON")
        return self._body


class FakeGraph:
    """Fake HTTP session: URL -> queue of responses, records calls."""

    def __init__(self) -> None:
        """Start with no routes and no recorded calls."""
        self.routes: dict[str, list[FakeResponse | Exception]] = {}
        self.calls: list[tuple[str, dict[str, str]]] = []

    def add(self, url: str, *responses: FakeResponse | Exception) -> None:
        """Queue responses for a URL; the last one repeats.

        Args:
            url: Exact URL the adapter should request.
            *responses: Responses (or exceptions to raise) in order.
        """
        self.routes.setdefault(url, []).extend(responses)

    def get(
        self, url: str, *, headers: Mapping[str, str], timeout: float
    ) -> FakeResponse:
        """Serve the next queued response for url.

        Args:
            url: Requested URL.
            headers: Request headers (recorded).
            timeout: Request timeout (must be set).

        Returns:
            The queued FakeResponse.

        Raises:
            AssertionError: URL was not expected.
            Exception: A queued exception, to simulate network errors.
        """
        assert timeout > 0
        self.calls.append((url, dict(headers)))
        queue = self.routes.get(url)
        if not queue:
            raise AssertionError(f"Unexpected Graph call: {url}")
        response = queue.pop(0) if len(queue) > 1 else queue[0]
        if isinstance(response, Exception):
            raise response
        return response

    def urls(self) -> list[str]:
        """Return requested URLs in order.

        Returns:
            List of URLs.
        """
        return [url for url, _ in self.calls]


def file_entry(
    item_id: str,
    name: str,
    etag: str = '"{v1},1"',
    size: int = 100,
) -> dict[str, Any]:
    """Build a Graph driveItem for a file.

    Args:
        item_id: Graph item ID.
        name: File name.
        etag: Item eTag.
        size: Size in bytes.

    Returns:
        driveItem JSON as a dict.
    """
    return {
        "id": item_id,
        "name": name,
        "eTag": etag,
        "size": size,
        "file": {"mimeType": "application/octet-stream"},
    }


def item_url(item_id: str) -> str:
    """Graph URL for one driveItem's metadata.

    Args:
        item_id: Graph item ID.

    Returns:
        URL string.
    """
    return f"{GRAPH_URL}/drives/{DRIVE_ID}/items/{item_id}"


def graph_error(status: int, code: str, message: str) -> FakeResponse:
    """Build a Graph-style error response.

    Args:
        status: HTTP status code.
        code: Graph error code.
        message: Graph error message.

    Returns:
        FakeResponse with a Graph error body.
    """
    return FakeResponse(status, {"error": {"code": code, "message": message}})


def synthetic_workbook() -> bytes:
    """Build a small, fixed workbook (sheet "Events", two rows).

    Returns:
        .xlsx bytes.
    """
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Events"
    ws.append(["event_id", "attendees"])
    ws.append(["E1", 10])
    ws.append(["E2", 12])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def config() -> SharePointConfig:
    """SharePoint location used by most tests."""
    return SharePointConfig(
        hostname="contoso.sharepoint.com",
        site_path="/sites/DataSite",
        library="Documents",
        folder="Incoming Files/2026",
    )


@pytest.fixture
def graph() -> FakeGraph:
    """Fake Graph with site and drive lookups already routed."""
    fake = FakeGraph()
    fake.add(SITE_URL, FakeResponse(body={"id": SITE_ID}))
    fake.add(
        DRIVES_URL,
        FakeResponse(
            body={
                "value": [
                    {"id": "b!other", "name": "Archive"},
                    {"id": DRIVE_ID, "name": "Documents"},
                ]
            }
        ),
    )
    return fake


@pytest.fixture
def sleeps() -> list[float]:
    """Records requested sleeps instead of sleeping."""
    return []


@pytest.fixture
def source(
    config: SharePointConfig, graph: FakeGraph, sleeps: list[float]
) -> SharePointSource:
    """SharePointSource wired to the fake Graph and a fake token."""
    return SharePointSource(
        config,
        token_provider=lambda: TOKEN,
        session=graph,
        sleep=sleeps.append,
    )


# --------------------------------------------------------------------
# Configuration
# --------------------------------------------------------------------


class TestSharePointConfig:
    """Location settings come from configuration; credentials don't."""

    def test_from_mapping_reads_location_settings(self) -> None:
        """All location settings are read from a config section."""
        cfg = SharePointConfig.from_mapping(
            {
                "hostname": "contoso.sharepoint.com",
                "site_path": "/sites/DataSite",
                "library": "Documents",
                "folder": "Incoming",
            }
        )
        assert cfg.hostname == "contoso.sharepoint.com"
        assert cfg.site_path == "/sites/DataSite"
        assert cfg.library == "Documents"
        assert cfg.folder == "Incoming"

    def test_folder_is_optional_and_defaults_to_library_root(self) -> None:
        """Omitting folder means the library root."""
        cfg = SharePointConfig.from_mapping(
            {
                "hostname": "contoso.sharepoint.com",
                "site_path": "/sites/DataSite",
                "library": "Documents",
            }
        )
        assert cfg.folder == ""

    def test_missing_required_setting_raises_config_error(self) -> None:
        """A missing required key is named in the error."""
        with pytest.raises(ConfigError, match="library"):
            SharePointConfig.from_mapping(
                {
                    "hostname": "contoso.sharepoint.com",
                    "site_path": "/sites/DataSite",
                }
            )

    def test_environment_placeholders_are_interpolated(
        self, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """${VAR} placeholders resolve from the environment.

        Args:
            monkeypatch: pytest fixture for environment variables.
        """
        monkeypatch.setenv("SP_HOST", "contoso.sharepoint.com")
        cfg = SharePointConfig.from_mapping(
            {
                "hostname": "${SP_HOST}",
                "site_path": "/sites/DataSite",
                "library": "Documents",
            }
        )
        assert cfg.hostname == "contoso.sharepoint.com"

    def test_unset_environment_placeholder_raises(self) -> None:
        """An unset ${VAR} fails loudly rather than being sent to Graph."""
        with pytest.raises(ConfigError, match="SP_UNSET_VAR"):
            SharePointConfig.from_mapping(
                {
                    "hostname": "${SP_UNSET_VAR}",
                    "site_path": "/sites/DataSite",
                    "library": "Documents",
                }
            )

    @pytest.mark.parametrize("key", ["client_secret", "tenant_id", "hostnme"])
    def test_unknown_settings_are_rejected(self, key: str) -> None:
        """Typos and credential fields are refused, not ignored.

        Args:
            key: An unexpected configuration key.
        """
        with pytest.raises(ConfigError, match=key):
            SharePointConfig.from_mapping(
                {
                    "hostname": "contoso.sharepoint.com",
                    "site_path": "/sites/DataSite",
                    "library": "Documents",
                    key: "x",
                }
            )

    @pytest.mark.parametrize(
        "hostname", ["https://contoso.sharepoint.com", "contoso/x", ""]
    )
    def test_hostname_must_be_bare_host(self, hostname: str) -> None:
        """Hostname is a host only: no scheme, no path.

        Args:
            hostname: An invalid hostname.
        """
        with pytest.raises(ConfigError, match="hostname"):
            SharePointConfig(hostname, "/sites/DataSite", "Documents")

    @pytest.mark.parametrize("site_path", ["sites/DataSite", "/", ""])
    def test_site_path_must_be_server_relative(self, site_path: str) -> None:
        """site_path starts with "/" and names a site.

        Args:
            site_path: An invalid site path.
        """
        with pytest.raises(ConfigError, match="site_path"):
            SharePointConfig("contoso.sharepoint.com", site_path, "Documents")

    def test_folder_slashes_are_normalised(self) -> None:
        """Leading/trailing slashes and backslashes are tidied."""
        cfg = SharePointConfig(
            "contoso.sharepoint.com",
            "/sites/DataSite",
            "Documents",
            folder="\\Incoming\\2026/",
        )
        assert cfg.folder == "Incoming/2026"

    def test_folder_cannot_climb_out_with_dotdot(self) -> None:
        """'..' segments are refused."""
        with pytest.raises(ConfigError, match="folder"):
            SharePointConfig(
                "contoso.sharepoint.com",
                "/sites/DataSite",
                "Documents",
                folder="Incoming/../Secret",
            )

    def test_extensions_default_to_excel_and_are_normalised(self) -> None:
        """Default is .xlsx/.xls; custom values are lower-cased."""
        default = SharePointConfig(
            "contoso.sharepoint.com", "/sites/DataSite", "Documents"
        )
        custom = SharePointConfig.from_mapping(
            {
                "hostname": "contoso.sharepoint.com",
                "site_path": "/sites/DataSite",
                "library": "Documents",
                "extensions": [".XLSX"],
            }
        )
        assert default.extensions == (".xlsx", ".xls")
        assert custom.extensions == (".xlsx",)

    def test_extension_without_leading_dot_is_rejected(self) -> None:
        """Extensions must start with a dot."""
        with pytest.raises(ConfigError, match="extensions"):
            SharePointConfig(
                "contoso.sharepoint.com",
                "/sites/DataSite",
                "Documents",
                extensions=("xlsx",),
            )


# --------------------------------------------------------------------
# Discovery
# --------------------------------------------------------------------


class TestSharePointSourceItems:
    """items() lists Excel files in the configured folder."""

    def test_items_returns_source_items_for_excel_files(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Graph driveItems map onto SourceItem fields.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(
                body={"value": [file_entry("ID1", "events.xlsx", size=42)]}
            ),
        )
        assert list(source.items()) == [
            SourceItem(
                name="events.xlsx",
                identity="ID1",
                version='"{v1},1"',
                size_bytes=42,
            )
        ]

    def test_items_resolves_site_then_drive_then_folder(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Lookups use the configured site, library and folder.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        list(source.items())
        assert graph.urls() == [SITE_URL, DRIVES_URL, CHILDREN_URL]

    def test_items_lists_library_root_when_no_folder(
        self, graph: FakeGraph
    ) -> None:
        """No folder configured means list the library root.

        Args:
            graph: Fake Graph fixture.
        """
        root_url = f"{GRAPH_URL}/drives/{DRIVE_ID}/root/children"
        graph.add(root_url, FakeResponse(body={"value": []}))
        cfg = SharePointConfig(
            "contoso.sharepoint.com", "/sites/DataSite", "Documents"
        )
        src = SharePointSource(cfg, lambda: TOKEN, session=graph)
        list(src.items())
        assert graph.urls()[-1] == root_url

    def test_items_skips_folders_and_non_excel_files(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Sub-folders and other file types are ignored.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(
                body={
                    "value": [
                        {"id": "F1", "name": "old", "folder": {}},
                        file_entry("ID1", "notes.docx"),
                        file_entry("ID2", "data.XLSX"),
                        file_entry("ID3", "legacy.xls"),
                    ]
                }
            ),
        )
        assert [i.identity for i in source.items()] == ["ID2", "ID3"]

    def test_items_follows_pagination(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Every @odata.nextLink page is fetched.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        page2 = f"{CHILDREN_URL}?$skiptoken=abc"
        page3 = f"{CHILDREN_URL}?$skiptoken=def"
        graph.add(
            CHILDREN_URL,
            FakeResponse(
                body={
                    "value": [file_entry("ID1", "a.xlsx")],
                    "@odata.nextLink": page2,
                }
            ),
        )
        graph.add(
            page2,
            FakeResponse(
                body={
                    "value": [file_entry("ID2", "b.xlsx")],
                    "@odata.nextLink": page3,
                }
            ),
        )
        graph.add(
            page3, FakeResponse(body={"value": [file_entry("ID3", "c.xlsx")]})
        )
        assert [i.identity for i in source.items()] == ["ID1", "ID2", "ID3"]

    def test_drive_lookup_follows_pagination(
        self, config: SharePointConfig
    ) -> None:
        """The library may be on a later page of the drives list.

        Args:
            config: SharePointConfig fixture.
        """
        fake = FakeGraph()
        fake.add(SITE_URL, FakeResponse(body={"id": SITE_ID}))
        page2 = f"{DRIVES_URL}?$skiptoken=x"
        fake.add(
            DRIVES_URL,
            FakeResponse(
                body={
                    "value": [{"id": "b!other", "name": "Archive"}],
                    "@odata.nextLink": page2,
                }
            ),
        )
        fake.add(
            page2,
            FakeResponse(
                body={"value": [{"id": DRIVE_ID, "name": "Documents"}]}
            ),
        )
        fake.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        list(SharePointSource(config, lambda: TOKEN, session=fake).items())
        assert fake.urls()[-1] == CHILDREN_URL

    def test_pagination_loop_is_detected(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """A nextLink that repeats stops with ServiceError, not forever.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(body={"value": [], "@odata.nextLink": CHILDREN_URL}),
        )
        with pytest.raises(ServiceError, match="pagination"):
            list(source.items())

    def test_items_are_sorted_by_name(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Order is stable across runs regardless of Graph ordering.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(
                body={
                    "value": [
                        file_entry("ID2", "b.xlsx"),
                        file_entry("ID3", "C.xlsx"),
                        file_entry("ID1", "a.xlsx"),
                    ]
                }
            ),
        )
        assert [i.name for i in source.items()] == [
            "a.xlsx",
            "b.xlsx",
            "C.xlsx",
        ]

    def test_site_and_drive_are_resolved_once(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Repeat listings reuse the resolved site and drive IDs.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        list(source.items())
        list(source.items())
        assert graph.urls().count(SITE_URL) == 1
        assert graph.urls().count(DRIVES_URL) == 1

    def test_unknown_library_raises_config_error(
        self, config: SharePointConfig
    ) -> None:
        """A library name that isn't on the site is a config error.

        Args:
            config: SharePointConfig fixture.
        """
        fake = FakeGraph()
        fake.add(SITE_URL, FakeResponse(body={"id": SITE_ID}))
        fake.add(
            DRIVES_URL,
            FakeResponse(body={"value": [{"id": "b!x", "name": "Archive"}]}),
        )
        src = SharePointSource(config, lambda: TOKEN, session=fake)
        with pytest.raises(ConfigError, match="Archive"):
            list(src.items())

    def test_library_name_match_ignores_case(
        self, config: SharePointConfig
    ) -> None:
        """SharePoint library names are case-insensitive.

        Args:
            config: SharePointConfig fixture.
        """
        fake = FakeGraph()
        fake.add(SITE_URL, FakeResponse(body={"id": SITE_ID}))
        fake.add(
            DRIVES_URL,
            FakeResponse(
                body={"value": [{"id": DRIVE_ID, "name": "DOCUMENTS"}]}
            ),
        )
        fake.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        list(SharePointSource(config, lambda: TOKEN, session=fake).items())

    def test_site_not_found_raises_config_error(
        self, config: SharePointConfig
    ) -> None:
        """A 404 on the site lookup points at the config.

        Args:
            config: SharePointConfig fixture.
        """
        fake = FakeGraph()
        fake.add(SITE_URL, graph_error(404, "itemNotFound", "nope"))
        src = SharePointSource(config, lambda: TOKEN, session=fake)
        with pytest.raises(ConfigError, match="site"):
            list(src.items())

    def test_folder_not_found_raises_config_error(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """A 404 on the folder listing points at the config.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, graph_error(404, "itemNotFound", "nope"))
        with pytest.raises(ConfigError, match="Incoming Files/2026"):
            list(source.items())

    def test_item_missing_required_field_raises_service_error(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """A driveItem without an eTag is refused, not guessed.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        entry = file_entry("ID1", "a.xlsx")
        del entry["eTag"]
        graph.add(CHILDREN_URL, FakeResponse(body={"value": [entry]}))
        with pytest.raises(ServiceError, match="eTag"):
            list(source.items())

    def test_non_json_response_raises_service_error(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """An HTML error page or garbage body is a ServiceError.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, FakeResponse(200, body=None))
        with pytest.raises(ServiceError, match="JSON"):
            list(source.items())


# --------------------------------------------------------------------
# Download
# --------------------------------------------------------------------


class TestSharePointSourceDownload:
    """download() returns the exact bytes of the discovered version."""

    @pytest.fixture
    def item(self) -> SourceItem:
        """A discovered file at version v1, 5 bytes."""
        return SourceItem(
            name="a.xlsx", identity="ID1", version='"{v1},1"', size_bytes=5
        )

    def test_download_returns_bytes_when_version_unchanged(
        self, source: SharePointSource, graph: FakeGraph, item: SourceItem
    ) -> None:
        """Unchanged file downloads normally.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            item: SourceItem fixture.
        """
        graph.add(item_url("ID1"), FakeResponse(body=file_entry("ID1", "a")))
        graph.add(f"{item_url('ID1')}/content", FakeResponse(content=b"hello"))
        assert source.download(item) == b"hello"

    def test_download_raises_version_mismatch_if_changed_before(
        self, source: SharePointSource, graph: FakeGraph, item: SourceItem
    ) -> None:
        """A newer version at download time is refused before fetching.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            item: SourceItem fixture.
        """
        graph.add(
            item_url("ID1"),
            FakeResponse(body=file_entry("ID1", "a", etag='"{v1},2"')),
        )
        with pytest.raises(VersionMismatch):
            source.download(item)
        assert f"{item_url('ID1')}/content" not in graph.urls()

    def test_download_raises_version_mismatch_if_changed_during(
        self, source: SharePointSource, graph: FakeGraph, item: SourceItem
    ) -> None:
        """A save that lands mid-download is caught afterwards.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            item: SourceItem fixture.
        """
        graph.add(
            item_url("ID1"),
            FakeResponse(body=file_entry("ID1", "a")),
            FakeResponse(body=file_entry("ID1", "a", etag='"{v1},2"')),
        )
        graph.add(f"{item_url('ID1')}/content", FakeResponse(content=b"hello"))
        with pytest.raises(VersionMismatch):
            source.download(item)

    def test_truncated_download_raises_service_error(
        self, source: SharePointSource, graph: FakeGraph, item: SourceItem
    ) -> None:
        """Byte count must match the size seen at discovery.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            item: SourceItem fixture.
        """
        graph.add(item_url("ID1"), FakeResponse(body=file_entry("ID1", "a")))
        graph.add(f"{item_url('ID1')}/content", FakeResponse(content=b"hel"))
        with pytest.raises(ServiceError, match="3 bytes"):
            source.download(item)

    def test_download_of_deleted_file_raises_not_found(
        self, source: SharePointSource, graph: FakeGraph, item: SourceItem
    ) -> None:
        """A file removed after discovery raises NotFound.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            item: SourceItem fixture.
        """
        graph.add(item_url("ID1"), graph_error(404, "itemNotFound", "gone"))
        with pytest.raises(NotFound):
            source.download(item)


# --------------------------------------------------------------------
# Authentication boundary, errors and retries
# --------------------------------------------------------------------


class TestSharePointSourceErrors:
    """HTTP failures are retried where sensible and reported clearly."""

    def test_bearer_token_from_provider_is_sent(
        self, source: SharePointSource, graph: FakeGraph
    ) -> None:
        """Each request carries the provider's token.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        list(source.items())
        assert all(
            headers["Authorization"] == f"Bearer {TOKEN}"
            for _, headers in graph.calls
        )

    def test_token_is_requested_per_request(
        self, config: SharePointConfig, graph: FakeGraph
    ) -> None:
        """The provider is asked every time, so it controls caching
        and refresh.

        Args:
            config: SharePointConfig fixture.
            graph: Fake Graph fixture.
        """
        graph.add(CHILDREN_URL, FakeResponse(body={"value": []}))
        provider = Mock(return_value=TOKEN)
        list(SharePointSource(config, provider, session=graph).items())
        assert provider.call_count == len(graph.calls)

    def test_token_provider_failure_raises_service_error(
        self, config: SharePointConfig, graph: FakeGraph
    ) -> None:
        """A failing provider becomes a chained ServiceError.

        Args:
            config: SharePointConfig fixture.
            graph: Fake Graph fixture.
        """
        provider = Mock(side_effect=RuntimeError("no credential"))
        src = SharePointSource(config, provider, session=graph)
        with pytest.raises(ServiceError, match="access token") as exc:
            list(src.items())
        assert isinstance(exc.value.__cause__, RuntimeError)
        assert graph.calls == []

    def test_empty_token_raises_service_error(
        self, config: SharePointConfig, graph: FakeGraph
    ) -> None:
        """An empty token is refused before any request.

        Args:
            config: SharePointConfig fixture.
            graph: Fake Graph fixture.
        """
        src = SharePointSource(config, lambda: "", session=graph)
        with pytest.raises(ServiceError, match="access token"):
            list(src.items())
        assert graph.calls == []

    @pytest.mark.parametrize("status", [401, 403])
    def test_auth_errors_raise_service_error_without_retry(
        self,
        source: SharePointSource,
        graph: FakeGraph,
        sleeps: list[float],
        status: int,
    ) -> None:
        """401/403 are not retried and name the Graph error code.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
            status: HTTP status under test.
        """
        graph.routes[SITE_URL] = [
            graph_error(status, "accessDenied", "Access denied")
        ]
        with pytest.raises(ServiceError, match=f"{status}.*accessDenied"):
            list(source.items())
        assert sleeps == []
        assert graph.urls() == [SITE_URL]

    def test_token_never_appears_in_errors_or_logs(
        self,
        source: SharePointSource,
        graph: FakeGraph,
        caplog: pytest.LogCaptureFixture,
    ) -> None:
        """The bearer token is not written to logs or messages.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            caplog: pytest log capture.
        """
        graph.add(
            CHILDREN_URL,
            graph_error(503, "serviceNotAvailable", "busy"),
            graph_error(401, "InvalidAuthenticationToken", "expired"),
        )
        with caplog.at_level(logging.DEBUG):
            with pytest.raises(ServiceError) as exc:
                list(source.items())
        assert TOKEN not in caplog.text
        assert TOKEN not in str(exc.value)

    def test_throttling_honours_retry_after(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """429 waits for Retry-After seconds, then succeeds.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(429, body={}, headers={"Retry-After": "7"}),
            FakeResponse(body={"value": [file_entry("ID1", "a.xlsx")]}),
        )
        assert len(list(source.items())) == 1
        assert sleeps == [7.0]

    def test_server_errors_back_off_exponentially(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """5xx without Retry-After waits 1s, 2s, ... then succeeds.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(
            CHILDREN_URL,
            graph_error(503, "serviceNotAvailable", "busy"),
            graph_error(502, "badGateway", "busy"),
            FakeResponse(body={"value": []}),
        )
        list(source.items())
        assert sleeps == [1.0, 2.0]

    def test_retry_after_is_capped(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """A huge Retry-After can't stall the job indefinitely.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(
            CHILDREN_URL,
            FakeResponse(429, body={}, headers={"Retry-After": "86400"}),
            FakeResponse(body={"value": []}),
        )
        list(source.items())
        assert sleeps == [60.0]

    def test_persistent_server_error_gives_up_with_service_error(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """After max attempts the last error is raised.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(CHILDREN_URL, graph_error(503, "serviceNotAvailable", "x"))
        with pytest.raises(ServiceError, match="503"):
            list(source.items())
        assert sleeps == [1.0, 2.0, 4.0]
        assert graph.urls().count(CHILDREN_URL) == 4

    def test_network_errors_are_retried_then_reported(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """Connection failures retry, then raise a chained ServiceError.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(CHILDREN_URL, requests.ConnectionError("reset"))
        with pytest.raises(ServiceError, match="ConnectionError") as exc:
            list(source.items())
        assert isinstance(exc.value.__cause__, requests.ConnectionError)
        assert len(sleeps) == 3

    def test_other_client_errors_are_not_retried(
        self, source: SharePointSource, graph: FakeGraph, sleeps: list[float]
    ) -> None:
        """A 400 is a bug or bad input; retrying won't help.

        Args:
            source: SharePointSource fixture.
            graph: Fake Graph fixture.
            sleeps: Recorded sleeps.
        """
        graph.add(CHILDREN_URL, graph_error(400, "invalidRequest", "bad"))
        with pytest.raises(ServiceError, match="400"):
            list(source.items())
        assert sleeps == []

    def test_max_attempts_must_be_positive(
        self, config: SharePointConfig, graph: FakeGraph
    ) -> None:
        """max_attempts < 1 is refused at construction.

        Args:
            config: SharePointConfig fixture.
            graph: Fake Graph fixture.
        """
        with pytest.raises(ValueError, match="max_attempts"):
            SharePointSource(
                config, lambda: TOKEN, session=graph, max_attempts=0
            )


# --------------------------------------------------------------------
# Fit with the existing pipeline
# --------------------------------------------------------------------


class TestSharePointSourceInPipeline:
    """SharePointSource drops into Pipeline in place of LocalSource."""

    def test_pipeline_loads_file_discovered_on_sharepoint(
        self, config: SharePointConfig, graph: FakeGraph
    ) -> None:
        """End to end with fake Graph, mock store and warehouse.

        Args:
            config: SharePointConfig fixture.
            graph: Fake Graph fixture.
        """
        workbook = synthetic_workbook()
        entry = file_entry("ID1", "events.xlsx", size=len(workbook))
        graph.add(CHILDREN_URL, FakeResponse(body={"value": [entry]}))
        graph.add(item_url("ID1"), FakeResponse(body=entry))
        graph.add(f"{item_url('ID1')}/content", FakeResponse(content=workbook))

        source: Source = SharePointSource(config, lambda: TOKEN, session=graph)
        store = MagicMock()
        store.get.return_value = None
        warehouse = Mock()
        warehouse.identity = "project.dataset.events"
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Events",
                    columns=[
                        Column(name="event_id", data_type="string"),
                        Column(name="attendees", data_type="integer"),
                    ],
                )
            ]
        )

        results = Pipeline(source, store, warehouse, contract).run()

        assert [r.outcome for r in results] == [Outcome.LOADED]
        assert results[0].rows_processed == 2
        store.put.assert_any_call(
            f"raw/ID1/{digest(workbook)}/events.xlsx", workbook
        )
