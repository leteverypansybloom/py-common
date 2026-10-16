"""Tests for workbook contract validation."""

from datetime import datetime
from io import BytesIO
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from py_common.contract import Column, Contract, Worksheet, read_sheet
from py_common.errors import ContractError, ValidationError
from py_common.fixtures import make_fixtures
import io
from unittest.mock import MagicMock
from unittest.mock import Mock
from unittest.mock import patch
from py_common.adapters.bigquery import BigQueryWarehouse
from py_common.contract import SheetRows
from py_common.model import Outcome
from py_common.model import SourceItem
from py_common.parquet import to_parquet
from py_common.pipeline import Pipeline


def _build_workbook(
    sheets: dict[str, tuple[list[str], list[list[Any]]]],
) -> bytes:
    """Build minimal .xlsx bytes with one or more named worksheets.

    Args:
        sheets: Maps worksheet name to (headers, data rows). The
            first entry becomes the workbook's active sheet, so
            tests can put a *different* sheet first than the
            contract expects, to prove active-sheet order is
            irrelevant to validation.

    Returns:
        Raw .xlsx file bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(title=name)
        for col_idx, header in enumerate(headers, start=1):
            ws.cell(row=1, column=col_idx, value=header)
        for row_idx, row in enumerate(rows, start=2):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _build_grid_workbook(sheets: dict[str, list[list[Any]]]) -> bytes:
    """Build .xlsx bytes, writing each sheet's rows exactly as given.

    Unlike _build_workbook, nothing is assumed about where the
    headings go, so tests can put title lines above them or leave
    cells blank.

    Args:
        sheets: Maps worksheet name to its rows, top to bottom.

    Returns:
        Raw .xlsx file bytes.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for row_idx, row in enumerate(rows, start=1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_valid_fixture(temp_dir, events_contract):
    """Valid fixture passes all validations."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (temp_dir / "fixtures" / "events_valid.xlsx").read_bytes()

    rows = events_contract.validate(excel_bytes, "events_valid.xlsx")

    assert rows == 2


def test_missing_worksheet(temp_dir, events_contract):
    """Missing worksheet raises ContractError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_missing_sheet.xlsx"
    ).read_bytes()

    with pytest.raises(ContractError, match="Missing worksheet"):
        events_contract.validate(
            excel_bytes,
            "events_missing_sheet.xlsx",
        )


def test_missing_column(temp_dir, events_contract):
    """Missing required column raises ContractError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_missing_column.xlsx"
    ).read_bytes()

    with pytest.raises(ContractError, match="Missing column"):
        events_contract.validate(
            excel_bytes,
            "events_missing_column.xlsx",
        )


def test_extra_column(temp_dir, events_contract):
    """Extra unexpected column raises ContractError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_extra_column.xlsx"
    ).read_bytes()

    with pytest.raises(ContractError, match="Unexpected columns"):
        events_contract.validate(
            excel_bytes,
            "events_extra_column.xlsx",
        )


def test_extra_columns_with_mixed_type_headers():
    """Multiple unexpected columns of different cell types don't crash.

    Regression test: 'unexpected' column names come straight from
    Excel header cells, which can hold different types (text,
    numbers, dates, ...). The old code passed those raw values
    straight to sorted(), which raises TypeError at runtime when
    asked to compare, e.g., a str header against an int header.
    Column names are stringified before comparing/sorting to avoid
    that crash and report a stable, readable error instead.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="event_id", data_type="string")],
            ),
        ]
    )
    excel_bytes = _build_workbook(
        {
            "Events": (
                ["event_id", "notes", 42],
                [["EVT001", "text header", "numeric header"]],
            ),
        }
    )

    with pytest.raises(ContractError, match="Unexpected columns"):
        contract.validate(excel_bytes, "mixed_headers.xlsx")


def test_invalid_data_type(temp_dir, events_contract):
    """Invalid data type raises ValidationError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_invalid_type.xlsx"
    ).read_bytes()

    with pytest.raises(ValidationError, match="expected integer"):
        events_contract.validate(
            excel_bytes,
            "events_invalid_type.xlsx",
        )


def test_missing_required_value(temp_dir, events_contract):
    """Null in required column raises ValidationError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_missing_value.xlsx"
    ).read_bytes()

    with pytest.raises(ValidationError, match="is required"):
        events_contract.validate(
            excel_bytes,
            "events_missing_value.xlsx",
        )


def test_duplicate_in_unique_column(temp_dir, events_contract):
    """Duplicate in unique column raises ValidationError."""
    make_fixtures(temp_dir / "fixtures")
    excel_bytes = (
        temp_dir / "fixtures" / "events_duplicate.xlsx"
    ).read_bytes()

    with pytest.raises(ValidationError, match="Duplicate value"):
        events_contract.validate(
            excel_bytes,
            "events_duplicate.xlsx",
        )


def test_malformed_excel(temp_dir, events_contract):
    """Malformed Excel file raises ContractError."""
    bad_excel = b"not a zip file"

    with pytest.raises(ContractError, match="Cannot parse Excel"):
        events_contract.validate(bad_excel, "bad.xlsx")


def test_empty_worksheets_raises_contract_error():
    """A Contract with no configured worksheets is a config error.

    Regression test: the old code did `self.worksheets[0]` and would
    raise an opaque IndexError instead of a clear ContractError.
    """
    contract = Contract(worksheets=[])
    excel_bytes = _build_workbook({"Sheet1": (["a"], [[1]])})

    with pytest.raises(ContractError, match="no worksheets configured"):
        contract.validate(excel_bytes, "empty.xlsx")


def test_multiple_worksheets_validated_against_matching_definitions():
    """Every configured worksheet is checked against its own schema.

    Fix for issue #5: the old code only ever validated
    self.worksheets[0]/wb.active, so a second worksheet's columns
    were never checked at all.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="event_id", data_type="string")],
            ),
            Worksheet(
                name="Employers",
                columns=[Column(name="employer_id", data_type="string")],
            ),
        ]
    )
    excel_bytes = _build_workbook(
        {
            "Events": (["event_id"], [["EVT001"]]),
            # Wrong column name in the second sheet - must be caught.
            "Employers": (["wrong_column"], [["EMP001"]]),
        }
    )

    with pytest.raises(ContractError, match="Missing column 'employer_id'"):
        contract.validate(excel_bytes, "multi.xlsx")


def test_active_sheet_order_has_no_effect():
    """Validation targets worksheets by name, never wb.active.

    Fix for issue #5: puts a non-contract sheet first (making it
    openpyxl's active sheet by default) and the contract's primary
    sheet second, then confirms the *contract's* first worksheet
    ("Events") is what gets validated and counted - not whichever
    sheet happens to be active in the workbook.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="event_id", data_type="string")],
            ),
        ]
    )
    excel_bytes = _build_workbook(
        {
            # "Summary" is first, so openpyxl makes it wb.active.
            "Summary": (["note"], [["irrelevant"]]),
            "Events": (
                ["event_id"],
                [["EVT001"], ["EVT002"], ["EVT003"]],
            ),
        }
    )

    rows = contract.validate(excel_bytes, "order.xlsx")

    assert rows == 3


def test_string_type_rejects_non_string_value():
    """A non-string value in a 'string' column is a ValidationError.

    Fix for issue #6: Contract.validate() previously only checked
    integer/float columns, so 'string' silently accepted anything.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="event_id", data_type="string")],
            ),
        ]
    )
    excel_bytes = _build_workbook({"Events": (["event_id"], [[12345]])})

    with pytest.raises(ValidationError, match="expected text"):
        contract.validate(excel_bytes, "bad_string.xlsx")


def test_date_type_validates_value():
    """Date columns are validated: wrong type rejected, real date ok.

    Fix for issue #6: 'date' was never checked before.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="held_on", data_type="date")],
            ),
        ]
    )

    bad_bytes = _build_workbook({"Events": (["held_on"], [["not-a-date"]])})
    with pytest.raises(ValidationError, match="expected date"):
        contract.validate(bad_bytes, "bad_date.xlsx")

    good_bytes = _build_workbook(
        {"Events": (["held_on"], [[datetime(2026, 1, 1)]])}
    )
    rows = contract.validate(good_bytes, "good_date.xlsx")
    assert rows == 1


def test_bool_rejected_by_integer_column():
    """A Python bool does not satisfy an 'integer' column.

    Fix for issue #6: bool is a subclass of int in Python, so
    isinstance(True, int) is True. Without an explicit exclusion,
    True/False would silently pass an integer field.
    """
    contract = Contract(
        worksheets=[
            Worksheet(
                name="Events",
                columns=[Column(name="attendees", data_type="integer")],
            ),
        ]
    )
    excel_bytes = _build_workbook({"Events": (["attendees"], [[True]])})

    with pytest.raises(ValidationError, match="expected integer"):
        contract.validate(excel_bytes, "bool.xlsx")


class TestWorksheetHeaderRow:
    """Worksheet.header_row says which row holds the headings."""

    def test_header_row_defaults_to_first_row(self):
        """Existing contracts keep reading headings from row 1."""
        assert Worksheet(name="Events").header_row == 1

    def test_header_row_below_one_is_rejected(self):
        """A header_row of 0 is a configuration mistake."""
        with pytest.raises(ValueError, match="header_row"):
            Worksheet(name="Events", header_row=0)


class TestContractValidateHeaderRow:
    """validate() reads headings from the configured row."""

    @pytest.fixture
    def contract(self) -> Contract:
        """Contract whose headings sit on row 2, under a title line."""
        return Contract(
            worksheets=[
                Worksheet(
                    name="Account signups",
                    header_row=2,
                    columns=[
                        Column(name="Org", data_type="string"),
                        Column(name="Size", data_type="string"),
                    ],
                )
            ]
        )

    def test_headings_on_row_two_are_found(self, contract):
        """A "Last updated" line above the headings is skipped.

        Args:
            contract: Contract fixture with header_row=2.
        """
        excel_bytes = _build_grid_workbook(
            {
                "Account signups": [
                    ["Last updated: 02/03/2026"],
                    ["Org", "Size"],
                    ["Org 1", "1000+"],
                    ["Org 2", "10-49"],
                ]
            }
        )

        rows = contract.validate(excel_bytes, "signups.xlsx")

        assert rows == 2

    def test_errors_report_excel_row_numbers(self, contract):
        """Row numbers in errors match what Excel shows.

        Args:
            contract: Contract fixture with header_row=2.
        """
        excel_bytes = _build_grid_workbook(
            {
                "Account signups": [
                    ["Last updated: 02/03/2026"],
                    ["Org", "Size"],
                    ["Org 1", 1000],
                ]
            }
        )

        with pytest.raises(ValidationError, match="Row 3: 'Size'"):
            contract.validate(excel_bytes, "signups.xlsx")

    def test_wrong_header_row_reports_missing_columns(self, contract):
        """Headings on row 1 don't satisfy a header_row=2 contract.

        Args:
            contract: Contract fixture with header_row=2.
        """
        excel_bytes = _build_grid_workbook(
            {"Account signups": [["Org", "Size"], ["Org 1", "1000+"]]}
        )

        with pytest.raises(ContractError, match="Missing column 'Org'"):
            contract.validate(excel_bytes, "signups.xlsx")


def _read(rows: list[list[Any]], header_row: int = 1):
    """Run read_sheet() over a one-sheet workbook built from rows.

    Args:
        rows: Sheet rows, top to bottom.
        header_row: Row holding the headings.

    Returns:
        read_sheet()'s (names, rows) result.
    """
    excel_bytes = _build_grid_workbook({"Sheet": rows})
    ws = load_workbook(BytesIO(excel_bytes), data_only=True)["Sheet"]
    return read_sheet(ws, header_row)


class TestReadSheet:
    """read_sheet() gives every column a present, unique name."""

    def test_blank_heading_over_data_is_named_by_letter(self):
        """Column B with no heading but with data becomes col_B."""
        names, rows = _read([["Count", None, "Org"], [1, "Nov", "Org 1"]])

        assert names == ["Count", "col_B", "Org"]
        assert rows == [(2, (1, "Nov", "Org 1"))]

    def test_blank_heading_over_empty_column_is_ignored(self):
        """A blank, empty column is only formatting, so it's dropped."""
        names, rows = _read([["Org", None, "Size"], ["Org 1", None, "S"]])

        assert names == ["Org", "Size"]
        assert rows == [(2, ("Org 1", "S"))]

    def test_repeated_headings_get_numbered(self):
        """The first "Other" keeps its name; later ones get _2, _3."""
        names, _ = _read([["Q1", "Other", "Q2", "Other", "Other"], [1] * 5])

        assert names == ["Q1", "Other", "Q2", "Other_2", "Other_3"]

    def test_renamed_repeat_skips_a_name_the_sheet_uses(self):
        """A repeat can't take "Other_2" if a real heading uses it."""
        names, _ = _read([["Other", "Other", "Other_2"], [1, 2, 3]])

        assert names == ["Other", "Other_3", "Other_2"]

    def test_headings_are_kept_exactly_as_written(self):
        """No trimming: contracts must match the sheet exactly."""
        names, _ = _read([["Organisation Industry ", 2026], ["x", "y"]])

        assert names == ["Organisation Industry ", "2026"]

    def test_empty_rows_are_skipped_and_row_numbers_kept(self):
        """Blank rows vanish; remaining rows keep their Excel numbers."""
        names, rows = _read(
            [["Org"], ["Org 1"], [None], ["   "], ["Org 2"]],
        )

        assert names == ["Org"]
        assert rows == [(2, ("Org 1",)), (5, ("Org 2",))]


class TestContractValidateHeadingNames:
    """validate() checks the names read_sheet() gives each column."""

    def test_repeated_heading_can_be_listed_by_its_new_name(self):
        """Contracts refer to the second "Other" as "Other_2"."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Survey info",
                    columns=[
                        Column(name="Other", data_type="string"),
                        Column(name="Other_2", data_type="string"),
                    ],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Survey info": [["Other", "Other"], ["a", "b"]]}
        )

        assert contract.validate(excel_bytes, "survey.xlsx") == 1

    def test_unlisted_blank_heading_column_is_unexpected(self):
        """A data column with no heading is reported, not ignored."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Account signups",
                    columns=[Column(name="Org", data_type="string")],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Account signups": [["Org", None], ["Org 1", "Nov"]]}
        )

        with pytest.raises(ContractError, match="Unexpected columns: col_B"):
            contract.validate(excel_bytes, "signups.xlsx")

    def test_trailing_blank_rows_are_not_counted(self):
        """Blank rows below the data don't count or fail checks."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Events",
                    columns=[Column(name="event_id", data_type="string")],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Events": [["event_id"], ["EVT001"], [None], [None]]}
        )

        assert contract.validate(excel_bytes, "events.xlsx") == 1


class TestWorksheetColumnSelection:
    """Worksheet checks its column selection settings up front."""

    def test_defaults_load_only_listed_columns(self):
        """Existing contracts keep loading just their listed columns."""
        ws_def = Worksheet(name="Events")

        assert ws_def.column_selection == "listed"
        assert ws_def.excluded == []

    def test_unknown_selection_is_rejected(self):
        """A typo in column_selection fails when the contract is built."""
        with pytest.raises(ValueError, match="column_selection"):
            Worksheet(name="Events", column_selection="everything")

    def test_column_both_listed_and_excluded_is_rejected(self):
        """A column can't be loaded and excluded at the same time."""
        with pytest.raises(ValueError, match="both listed and excluded"):
            Worksheet(
                name="Events",
                columns=[Column(name="Email", data_type="string")],
                excluded=["Email"],
            )


class TestContractValidateColumnSelection:
    """validate() accepts excluded and, in all_except, new columns."""

    def test_excluded_column_is_not_unexpected(self):
        """A known personal-data column can stay in the sheet."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    columns=[
                        Column(name="Participant ID", data_type="string")
                    ],
                    excluded=["Email Address"],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID", "Email Address"], ["0001", "a@b"]]}
        )

        assert contract.validate(excel_bytes, "peers.xlsx") == 1

    def test_all_except_accepts_new_columns(self):
        """A column added to the sheet doesn't stop the load."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    columns=[
                        Column(name="Participant ID", data_type="string")
                    ],
                    column_selection="all_except",
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID", "New question"], ["0001", "yes"]]}
        )

        assert contract.validate(excel_bytes, "peers.xlsx") == 1

    def test_all_except_still_requires_listed_columns(self):
        """Listed columns must be present in either mode."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    columns=[
                        Column(name="Participant ID", data_type="string")
                    ],
                    column_selection="all_except",
                )
            ]
        )
        excel_bytes = _build_grid_workbook({"Peers": [["ID"], ["0001"]]})

        with pytest.raises(ContractError, match="Missing column"):
            contract.validate(excel_bytes, "peers.xlsx")


class TestContractExtract:
    """extract() returns only the columns the contract selects."""

    def test_listed_mode_returns_listed_columns_in_contract_order(self):
        """Unlisted and excluded columns are left out."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    columns=[
                        Column(name="Sector", data_type="string"),
                        Column(name="Participant ID", data_type="string"),
                    ],
                    excluded=["Email Address"],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {
                "Peers": [
                    ["Participant ID", "Email Address", "Sector"],
                    ["0001", "a@b", "Public"],
                ]
            }
        )

        sheet = contract.extract(excel_bytes)["Peers"]

        assert sheet.columns == ["Sector", "Participant ID"]
        assert sheet.rows == [(2, ("Public", "0001"))]

    def test_all_except_mode_drops_only_excluded_columns(self):
        """Everything else is kept, in the sheet's order."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    column_selection="all_except",
                    excluded=["First Name", "Email Address"],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {
                "Peers": [
                    [
                        "Participant ID",
                        "First Name",
                        "Sector",
                        "Email Address",
                    ],
                    ["0001", "Jan", "Public", "a@b"],
                ]
            }
        )

        sheet = contract.extract(excel_bytes)["Peers"]

        assert sheet.columns == ["Participant ID", "Sector"]
        assert sheet.rows == [(2, ("0001", "Public"))]

    def test_listed_column_missing_from_sheet_is_none(self):
        """A missing listed column comes back empty, not as an error."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    columns=[
                        Column(name="Participant ID", data_type="string"),
                        Column(name="Sector", data_type="string"),
                    ],
                )
            ]
        )
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID"], ["0001"]]}
        )

        sheet = contract.extract(excel_bytes)["Peers"]

        assert sheet.rows == [(2, ("0001", None))]

    def test_every_configured_sheet_is_returned_in_order(self):
        """One SheetRows per worksheet, using each sheet's header_row."""
        contract = Contract(
            worksheets=[
                Worksheet(
                    name="Survey info",
                    header_row=2,
                    columns=[Column(name="Org", data_type="string")],
                ),
                Worksheet(
                    name="Account signups",
                    columns=[Column(name="Org", data_type="string")],
                ),
            ]
        )
        excel_bytes = _build_grid_workbook(
            {
                "Account signups": [["Org"], ["Org 2"]],
                "Survey info": [["Last updated"], ["Org"], ["Org 1"]],
            }
        )

        sheets = contract.extract(excel_bytes)

        assert list(sheets) == ["Survey info", "Account signups"]
        assert sheets["Survey info"].rows == [(3, ("Org 1",))]
        assert sheets["Account signups"].rows == [(2, ("Org 2",))]


class TestWorksheetOnMismatch:
    """Worksheet.on_mismatch defaults to today's quarantine."""

    def test_defaults_to_quarantine(self):
        """Existing contracts keep rejecting mismatched files."""
        assert Worksheet(name="Events").on_mismatch == "quarantine"

    def test_unknown_mode_is_rejected(self):
        """A typo in on_mismatch fails when the contract is built."""
        with pytest.raises(ValueError, match="on_mismatch"):
            Worksheet(name="Events", on_mismatch="ignore")


class TestContractValidateWarnMode:
    """In "warn" mode, problems are recorded and the file passes."""

    @staticmethod
    def _contract(on_mismatch: str) -> Contract:
        """Contract for a "Peers" sheet with an ID and a size.

        Args:
            on_mismatch: "quarantine" or "warn".

        Returns:
            Contract.
        """
        return Contract(
            worksheets=[
                Worksheet(
                    name="Peers",
                    on_mismatch=on_mismatch,
                    columns=[
                        Column(
                            name="Participant ID",
                            data_type="string",
                            unique=True,
                        ),
                        Column(name="Business Size", data_type="string"),
                    ],
                )
            ]
        )

    def test_renamed_column_is_a_warning(self):
        """A renamed heading is reported, and the rows still count."""
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID", "Size"], ["0001", "50-249"]]}
        )
        warnings: list[str] = []

        rows = self._contract("warn").validate(
            excel_bytes, "peers.xlsx", warnings=warnings
        )

        assert rows == 1
        assert len(warnings) == 1
        assert "Missing column 'Business Size'" in warnings[0]
        assert "Unexpected columns: Size" in warnings[0]

    def test_renamed_column_quarantines_in_quarantine_mode(self):
        """The same file is rejected when on_mismatch is quarantine."""
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID", "Size"], ["0001", "50-249"]]}
        )

        with pytest.raises(ContractError, match="Missing column"):
            self._contract("quarantine").validate(excel_bytes, "peers.xlsx")

    def test_missing_column_is_not_reported_on_every_row(self):
        """One header warning, not a "required" error per row."""
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID"], ["0001"], ["0002"]]}
        )
        warnings: list[str] = []

        self._contract("warn").validate(
            excel_bytes, "peers.xlsx", warnings=warnings
        )

        assert len(warnings) == 1
        assert "is required" not in warnings[0]

    def test_row_problems_are_warnings(self):
        """A duplicate and a wrong type don't stop the file."""
        excel_bytes = _build_grid_workbook(
            {
                "Peers": [
                    ["Participant ID", "Business Size"],
                    ["0001", "50-249"],
                    ["0001", 250],
                ]
            }
        )
        warnings: list[str] = []

        rows = self._contract("warn").validate(
            excel_bytes, "peers.xlsx", warnings=warnings
        )

        assert rows == 2
        assert "Row 3: Duplicate value" in warnings[0]
        assert "Row 3: 'Business Size' is int" in warnings[0]

    def test_row_problems_quarantine_in_quarantine_mode(self):
        """The same rows are rejected when on_mismatch is quarantine."""
        excel_bytes = _build_grid_workbook(
            {
                "Peers": [
                    ["Participant ID", "Business Size"],
                    ["0001", "50-249"],
                    ["0001", "50-249"],
                ]
            }
        )

        with pytest.raises(ValidationError, match="Duplicate value"):
            self._contract("quarantine").validate(excel_bytes, "peers.xlsx")

    def test_warnings_are_tolerated_without_a_list(self):
        """Callers that pass no list still get a passing file."""
        excel_bytes = _build_grid_workbook(
            {"Peers": [["Participant ID"], ["0001"]]}
        )

        assert self._contract("warn").validate(excel_bytes, "p.xlsx") == 1

    def test_valid_file_adds_no_warnings(self):
        """Nothing is recorded when the sheet matches."""
        excel_bytes = _build_grid_workbook(
            {
                "Peers": [
                    ["Participant ID", "Business Size"],
                    ["0001", "50-249"],
                ]
            }
        )
        warnings: list[str] = []

        self._contract("warn").validate(
            excel_bytes, "peers.xlsx", warnings=warnings
        )

        assert warnings == []


ITEM = SourceItem("f.xlsx", "f.xlsx", "1", 1)


SECRET = "jo.bloggs@example.org"


def _book(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Events"
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestReservedColumn:
    """N02: a source `_loaded_at` column quarantines, not retries."""

    def _warehouse(self) -> BigQueryWarehouse:
        client = MagicMock()
        with patch("py_common.adapters.bigquery.bigquery.Client") as cls:
            cls.return_value = client
            warehouse = BigQueryWarehouse("p", "d", "t")
        from types import SimpleNamespace

        client.get_table.return_value = SimpleNamespace(
            schema=[
                SimpleNamespace(name="_loaded_at", field_type="STRING"),
            ]
        )
        client.load_table_from_file.return_value = Mock(
            job_id="l", output_rows=1
        )
        return warehouse

    def test_raises_contract_error(self) -> None:
        with pytest.raises(ContractError, match="_loaded_at"):
            self._warehouse().load("k", "c" * 64, b"pq", 1, ITEM)

    def test_pipeline_quarantines_the_file(self) -> None:
        objects: dict[str, bytes] = {}
        store = MagicMock()
        store.get.side_effect = objects.get
        store.put.side_effect = objects.__setitem__
        source = Mock()
        source.download.return_value = _book([["_loaded_at"], ["x"]])
        contract = Contract(
            worksheets=[Worksheet("Events", [Column("_loaded_at", "string")])]
        )
        pipeline = Pipeline(source, store, self._warehouse(), contract)

        result = pipeline.process(ITEM)

        assert result.outcome == Outcome.QUARANTINED
        assert any(k.startswith("quarantine/") for k in objects)


class TestNoValuesInMessages:
    """M13: errors and warnings name rows and columns, not values."""

    def test_type_error_has_no_value(self) -> None:
        contract = Contract(
            worksheets=[Worksheet("Events", [Column("n", "integer")])]
        )
        with pytest.raises(ValidationError) as info:
            contract.read(_book([["n"], [SECRET]]), "f.xlsx")
        assert SECRET not in str(info.value)

    def test_duplicate_key_message_has_no_value(self) -> None:
        contract = Contract(
            key_columns=["id"],
            worksheets=[Worksheet("Events", [Column("id", "string")])],
        )
        with pytest.raises(ValidationError, match="Row 3") as info:
            contract.read(_book([["id"], [SECRET], [SECRET]]), "f.xlsx")
        assert SECRET not in str(info.value)

    def test_conversion_warning_has_no_value(self) -> None:
        warnings: list[str] = []
        sheet = Worksheet(
            "Events", [Column("n", "integer")], on_mismatch="warn"
        )
        to_parquet(SheetRows(["n"], [(2, (SECRET,))]), sheet, "f", warnings)
        assert warnings and SECRET not in warnings[0]
