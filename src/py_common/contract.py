"""Workbook contract validation.

A contract specifies required worksheets, columns, data types,
and business rules. Validation is deterministic and repeatable.
"""

from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet as XlWorksheet

from py_common.errors import ContractError, ValidationError

# Maps a Column.data_type to the Python type(s) a cell value must be
# an instance of, and the word used in error messages. bool is
# excluded everywhere because Python's bool is a subclass of int and
# would otherwise silently pass "integer"/"float" columns.
_TYPE_CHECKS: dict[str, tuple[type | tuple[type, ...], str]] = {
    "integer": (int, "integer"),
    "float": ((int, float), "number"),
    "string": (str, "text"),
    # datetime.datetime is a subclass of datetime.date, so this
    # covers both date-only and date+time cell values.
    "date": (date, "date"),
}


@dataclass
class Column:
    """Column definition within a worksheet.

    Attributes:
        name: Column header (must match worksheet exactly).
        data_type: Expected type: 'string', 'integer', 'float', 'date'.
        nullable: Whether column may contain null/empty values.
        unique: Whether all values must be distinct (no duplicates).
    """

    name: str
    data_type: Literal["string", "integer", "float", "date"]
    nullable: bool = False
    unique: bool = False


@dataclass
class Worksheet:
    """Worksheet definition within a workbook.

    Attributes:
        name: Worksheet name (must match Excel sheet name).
        columns: List of required columns.
        header_row: Row holding the column headings (1 = first row).
            Data starts on the row after it. Sheets laid out for
            people often put a title or "Last updated" line above
            the headings.
        column_selection: Which columns are loaded. "listed" (the
            default) loads only `columns`, and any other column in
            the sheet is a mismatch unless it is in `excluded`.
            "all_except" loads every column except `excluded`, so
            new columns are picked up without a contract change.
        excluded: Columns known to be in the sheet that are never
            loaded, e.g. names or email addresses the analysis
            doesn't need. Names are as read_sheet() gives them.
        on_mismatch: What a mismatch does. "quarantine" (the default)
            rejects the file, as before. "warn" records the problem
            as a warning and loads the file anyway: missing columns
            load as empty, and row-level problems (wrong type,
            missing required value, duplicate) are warnings too, left
            for the cleaning layer to deal with row by row.
    """

    name: str
    columns: list[Column] = field(default_factory=list)
    header_row: int = 1
    column_selection: Literal["listed", "all_except"] = "listed"
    excluded: list[str] = field(default_factory=list)
    on_mismatch: Literal["quarantine", "warn"] = "quarantine"

    def __post_init__(self) -> None:
        """Check settings that would otherwise fail confusingly later.

        Raises:
            ValueError: header_row is less than 1, column_selection
                or on_mismatch is unknown, or a column is both listed
                and excluded.
        """
        if self.on_mismatch not in ("quarantine", "warn"):
            raise ValueError(
                f"Worksheet '{self.name}': on_mismatch must be "
                f"'quarantine' or 'warn', got {self.on_mismatch!r}"
            )
        if self.header_row < 1:
            raise ValueError(
                f"Worksheet '{self.name}': header_row must be 1 or "
                f"more, got {self.header_row}"
            )
        if self.column_selection not in ("listed", "all_except"):
            raise ValueError(
                f"Worksheet '{self.name}': column_selection must be "
                f"'listed' or 'all_except', got "
                f"{self.column_selection!r}"
            )
        both = sorted({c.name for c in self.columns} & set(self.excluded))
        if both:
            raise ValueError(
                f"Worksheet '{self.name}': columns both listed and "
                f"excluded: {', '.join(both)}"
            )


@dataclass
class SheetRows:
    """The columns and rows to load from one worksheet.

    Attributes:
        columns: Names of the columns to load. In "listed" mode this
            is the contract's column order; in "all_except" mode it
            is the sheet's order.
        rows: (Excel row number, values lined up with `columns`) for
            each non-empty data row. A listed column missing from
            the sheet has None in every row.
    """

    columns: list[str]
    rows: list[tuple[int, tuple[Any, ...]]]


@dataclass
class Contract:
    """Workbook schema and validation rules.

    Attributes:
        key_columns: Columns of the primary (first) worksheet that
            identify a row, used to MERGE into the warehouse. Each
            must be a listed, non-nullable column, and together they
            must be unique within a file. Empty means append.
        worksheets: List of required worksheets.
    """

    key_columns: list[str] = field(default_factory=list)
    worksheets: list[Worksheet] = field(default_factory=list)

    def __post_init__(self) -> None:
        """Check key_columns can actually key a MERGE.

        A key the data doesn't have used to be dropped silently, and
        a NULL key never matches in a MERGE, so both turned upserts
        into duplicate inserts. Refuse them when the contract is
        built instead.

        Raises:
            ValueError: A key column isn't a listed column of the
                primary worksheet, is nullable, or there are key
                columns but no worksheets.
        """
        if not self.key_columns:
            return
        if not self.worksheets:
            raise ValueError(
                "Contract has key_columns but no worksheets to take "
                "them from"
            )
        primary = self.worksheets[0]
        listed = {column.name: column for column in primary.columns}
        for key in self.key_columns:
            column = listed.get(key)
            if column is None:
                raise ValueError(
                    f"Key column '{key}' is not a listed column of the "
                    f"primary worksheet '{primary.name}'"
                )
            if column.nullable:
                raise ValueError(
                    f"Key column '{key}' of worksheet '{primary.name}' "
                    "is nullable; key columns need nullable=False"
                )

    def validate(
        self,
        excel_bytes: bytes,
        source_name: str,
        warnings: list[str] | None = None,
    ) -> int:
        """Validate workbook against this contract.

        Every worksheet in `self.worksheets` is validated against its
        matching sheet, looked up by name — never the workbook's
        active sheet, which has no relation to what the contract
        expects and can silently differ from it.

        Args:
            excel_bytes: Raw Excel file bytes.
            source_name: Filename for error messages.
            warnings: List to append problems to for worksheets whose
                on_mismatch is "warn" (mutated). If None, those
                problems are still tolerated but not reported.

        Returns:
            Row count of the first configured worksheet if valid.

        Raises:
            ContractError: Contract is misconfigured, or worksheet
                structure doesn't match it.
            ValidationError: Data violates contract rules.
        """
        wb = self._open(excel_bytes, source_name)
        if warnings is None:
            warnings = []

        primary_rows = 0
        for index, ws_def in enumerate(self.worksheets):
            ws = wb[ws_def.name]
            # Keys belong to the primary worksheet (see __post_init__).
            keys = self.key_columns if index == 0 else []
            rows = self._validate_worksheet(
                ws, ws_def, source_name, warnings, keys
            )
            if index == 0:
                primary_rows = rows

        return primary_rows

    def extract(
        self,
        excel_bytes: bytes,
        source_name: str = "workbook",
    ) -> dict[str, SheetRows]:
        """Read the columns and rows to load from every worksheet.

        Applies each worksheet's header_row, heading names and column
        selection. Does not validate: call validate() first.

        Args:
            excel_bytes: Raw Excel file bytes.
            source_name: Filename for error messages.

        Returns:
            SheetRows for each configured worksheet, keyed by name,
            in contract order.

        Raises:
            ContractError: Contract is misconfigured, the workbook
                can't be parsed, or a worksheet is missing.
        """
        wb = self._open(excel_bytes, source_name)
        sheets: dict[str, SheetRows] = {}
        for ws_def in self.worksheets:
            names, rows = read_sheet(wb[ws_def.name], ws_def.header_row)
            index = {name: i for i, name in enumerate(names)}
            if ws_def.column_selection == "listed":
                selected = [col.name for col in ws_def.columns]
            else:
                skip = set(ws_def.excluded)
                selected = [name for name in names if name not in skip]
            sheets[ws_def.name] = SheetRows(
                columns=selected,
                rows=[
                    (
                        row_number,
                        tuple(
                            values[index[name]] if name in index else None
                            for name in selected
                        ),
                    )
                    for row_number, values in rows
                ],
            )
        return sheets

    def _open(self, excel_bytes: bytes, source_name: str) -> Workbook:
        """Parse the workbook and check every worksheet is present.

        Args:
            excel_bytes: Raw Excel file bytes.
            source_name: Filename for error messages.

        Returns:
            The openpyxl workbook.

        Raises:
            ContractError: Contract has no worksheets, the workbook
                can't be parsed, or a configured worksheet is missing.
        """
        if not self.worksheets:
            raise ContractError(
                f"{source_name}: Contract has no worksheets configured"
            )

        try:
            wb = load_workbook(
                # file-like object:
                filename=BytesIO(excel_bytes),
                data_only=True,
                read_only=False,
            )
        except Exception as e:
            raise ContractError(
                f"{source_name}: Cannot parse Excel: {e}"
            ) from e

        # Check every configured worksheet exists before reading any
        # of them, so one missing sheet reports cleanly instead of a
        # KeyError on the next lookup.
        missing = [
            ws_def.name
            for ws_def in self.worksheets
            if ws_def.name not in wb.sheetnames
        ]
        if missing:
            raise ContractError(
                f"{source_name}: Missing worksheet(s) "
                f"{', '.join(missing)}. Found: {', '.join(wb.sheetnames)}"
            )
        return wb

    def _validate_worksheet(
        self,
        ws: XlWorksheet,
        ws_def: Worksheet,
        source_name: str,
        warnings: list[str],
        key_columns: list[str] | None = None,
    ) -> int:
        """Validate one worksheet against its Worksheet definition.

        Args:
            ws: The openpyxl worksheet to validate.
            ws_def: Contract definition for this worksheet.
            source_name: Filename for error messages.
            warnings: Where problems go when ws_def.on_mismatch is
                "warn" (mutated).
            key_columns: Columns whose values together must be unique
                per row (the contract's merge key); empty for none.

        Returns:
            Number of non-empty data rows (below the heading row) if
            valid.

        Raises:
            ContractError: Column structure doesn't match ws_def, and
                ws_def.on_mismatch is "quarantine".
            ValidationError: Data violates ws_def's column rules, and
                ws_def.on_mismatch is "quarantine".
        """
        errors: list[str] = []
        header = f"{source_name} [{ws_def.name}]"

        headers, rows = read_sheet(ws, ws_def.header_row)
        for col_def in ws_def.columns:
            if col_def.name not in headers:
                errors.append(f"Missing column '{col_def.name}'")

        if ws_def.column_selection == "listed":
            # read_sheet() has already turned every heading into text
            # (headings can be numbers or dates too), so they compare
            # and sort cleanly here.
            expected = {col.name for col in ws_def.columns}
            unexpected = set(headers) - expected - set(ws_def.excluded)
            if unexpected:
                errors.append(
                    f"Unexpected columns: {', '.join(sorted(unexpected))}"
                )

        if errors:
            message = f"{header}: {'; '.join(errors)}"
            if ws_def.on_mismatch == "quarantine":
                raise ContractError(message)
            warnings.append(message)

        # Column position in the sheet may not match ws_def.columns'
        # order, so look up each column's index by its header name
        # rather than assuming the two lists are aligned.
        header_index = {h: i for i, h in enumerate(headers)}

        row_errors: list[str] = []
        seen_values: dict[str, set[Any]] = {
            col.name: set() for col in ws_def.columns if col.unique
        }

        # A missing column has already been reported once above (only
        # possible in "warn" mode); don't report it again on every row.
        present = [c for c in ws_def.columns if c.name in header_index]

        # The composite merge key must be unique, or the warehouse
        # MERGE fails ("must match at most one source row"). A key
        # column missing from the sheet was reported above.
        keys = [k for k in key_columns or [] if k in header_index]
        if len(keys) != len(key_columns or []):
            keys = []
        first_seen: dict[tuple[Any, ...], int] = {}

        for row_idx, row in rows:
            if keys:
                key_values = tuple(
                    (
                        row[header_index[k]]
                        if header_index[k] < len(row)
                        else None
                    )
                    for k in keys
                )
                # A NULL key is reported as "required" below.
                if not any(_is_blank(v) for v in key_values):
                    if key_values in first_seen:
                        row_errors.append(
                            f"Row {row_idx}: Duplicate key {key_values} "
                            f"in ({', '.join(keys)}), first on row "
                            f"{first_seen[key_values]}"
                        )
                    else:
                        first_seen[key_values] = row_idx

            for col_def in present:
                col_idx = header_index[col_def.name]
                value = row[col_idx] if col_idx < len(row) else None

                # Check nullable
                if value is None and not col_def.nullable:
                    row_errors.append(
                        f"Row {row_idx}: '{col_def.name}' is required"
                    )
                    continue

                if value is not None:
                    expected_type, label = _TYPE_CHECKS[col_def.data_type]
                    valid = isinstance(
                        value, expected_type
                    ) and not isinstance(value, bool)
                    if not valid:
                        row_errors.append(
                            f"Row {row_idx}: '{col_def.name}' "
                            f"is {type(value).__name__}, "
                            f"expected {label}"
                        )

                    # Check uniqueness
                    if col_def.unique:
                        if value in seen_values[col_def.name]:
                            row_errors.append(
                                f"Row {row_idx}: Duplicate value "
                                f"in unique column '{col_def.name}'"
                            )
                        seen_values[col_def.name].add(value)

        if row_errors:
            message = f"{header}: {'; '.join(row_errors[:5])}" + (
                f"... and {len(row_errors) - 5} more"
                if len(row_errors) > 5
                else ""
            )
            if ws_def.on_mismatch == "quarantine":
                raise ValidationError(message)
            warnings.append(message)

        return len(rows)


def read_sheet(
    ws: XlWorksheet, header_row: int
) -> tuple[list[str], list[tuple[int, tuple[Any, ...]]]]:
    """Read a worksheet's column names and the data rows below them.

    Headings are turned into names that are always present and
    unique, so no column is lost on the way to the warehouse:

    - a heading cell is used as text, exactly as written (no
      trimming, so contracts must match it exactly);
    - a blank heading over a column that holds data is named after
      its column letter, e.g. "col_B";
    - a blank heading over an empty column is ignored, as it is only
      formatting;
    - a repeated heading keeps its name the first time and gets
      "_2", "_3", ... after that, left to right, skipping any name
      the sheet already uses.

    Rows with no values at all are skipped.

    Args:
        ws: The openpyxl worksheet to read.
        header_row: Row holding the headings (1 = first row).

    Returns:
        (names, rows): the column names, and each data row as
        (Excel row number, values lined up with names). Row numbers
        are the ones shown in Excel, so errors can be traced back to
        the sheet.
    """
    grid = list(ws.iter_rows(min_row=header_row, values_only=True))
    if not grid:
        return [], []
    headings = list(grid[0])
    data = [
        (header_row + offset, values)
        for offset, values in enumerate(grid[1:], start=1)
        if not all(_is_blank(value) for value in values)
    ]

    # Reserve every heading actually written in the sheet, so a
    # renamed repeat ("Other" -> "Other_2") never takes a name that
    # a later heading uses itself.
    written = {str(h) for h in headings if not _is_blank(h)}
    names: list[str] = []
    keep: list[int] = []
    for index, heading in enumerate(headings):
        if _is_blank(heading):
            has_data = any(
                index < len(values) and not _is_blank(values[index])
                for _, values in data
            )
            if not has_data:
                continue
            base = f"col_{get_column_letter(index + 1)}"
        else:
            base = str(heading)
        name = base
        suffix = 2
        while name in names or (name != base and name in written):
            name = f"{base}_{suffix}"
            suffix += 1
        names.append(name)
        keep.append(index)

    rows = [
        (
            row_number,
            tuple(values[i] if i < len(values) else None for i in keep),
        )
        for row_number, values in data
    ]
    return names, rows


def _is_blank(value: Any) -> bool:
    """Whether a cell holds nothing (empty, or only whitespace).

    Args:
        value: Cell value from openpyxl.

    Returns:
        True for None or a whitespace-only string.
    """
    return value is None or (isinstance(value, str) and not value.strip())
